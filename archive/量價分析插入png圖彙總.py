import pandas as pd
from decimal import Decimal, ROUND_HALF_UP
from pandas import ExcelWriter
import matplotlib
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook
from openpyxl.drawing.image import Image
import io

# ====== 設置 matplotlib 正確顯示中文 ======
matplotlib.rcParams['font.sans-serif'] = ['Microsoft JhengHei']  # 使用微軟正黑體
matplotlib.rcParams['axes.unicode_minus'] = False  # 正負號正常顯示

# ====== 檔案路徑設定 ======
input_path  = r"C:\Users\\Desktop\00693U_0604.csv"
output_path = r"C:\Users\\Desktop\00693U_0604分析結果含表百分比.xlsx"

# ====== 1) 讀取原始檔案 ======
df = pd.read_csv(input_path, encoding="big5", dtype=str, usecols=[0, 1, 2, 3, 4])
df.columns = ["序號", "券商", "股價", "買進股數", "賣出股數"]

# ====== 2) 數字清理 ======
def clean_int(s: str) -> int:
    if not s: return 0
    s = s.replace(",", "")
    digits = "".join(ch for ch in s if ch.isdigit())
    return int(digits) if digits else 0

df["買進股數"] = df["買進股數"].map(clean_int).astype(int)
df["賣出股數"] = df["賣出股數"].map(clean_int).astype(int)

# ====== 3) 股價標準化 ======
def norm_px(s: str) -> Decimal:
    d = Decimal(s.strip())
    return d.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

df["股價"] = df["股價"].map(norm_px)

# ====== 4) 第二工作表：買賣價量與家數 ======
price_total = df.groupby("股價")[["買進股數","賣出股數"]].sum().reset_index()
broker_count = df.groupby("股價").agg(
    買進家數=("買進股數", lambda x: (x>0).sum()),
    賣出家數=("賣出股數", lambda x: (x>0).sum())
).reset_index()
price_summary = price_total.merge(broker_count, on="股價").sort_values("股價")

# ====== 5) 第三工作表：券商明細（含占比與標記） ======
price_broker = df.groupby(["股價","券商"])[["買進股數","賣出股數"]].sum().reset_index()
merged = price_broker.merge(price_total, on="股價", suffixes=("","_總"))

merged["買進占比"] = (merged["買進股數"] / merged["買進股數_總"].replace(0,1) * 100).round(3).astype(str) + "%"
merged["賣出占比"] = (merged["賣出股數"] / merged["賣出股數_總"].replace(0,1) * 100).round(3).astype(str) + "%"

def mark_top(group):
    group["最高買進"] = group["買進股數"] == group["買進股數"].max()
    group["最高賣出"] = group["賣出股數"] == group["賣出股數"].max()
    return group
merged = merged.groupby("股價", group_keys=False).apply(mark_top)

# ====== 在繪圖前補上最高買進/賣出股數欄位 ======
# 最高買進：每個股價取券商買進股數的最大值
top_buy = (
    merged.loc[merged["最高買進"], ["股價", "買進股數"]]
    .groupby("股價")["買進股數"]
    .max()
)
price_summary["最高買進股數"] = price_summary["股價"].map(top_buy).fillna(0)

# 最高賣出：每個股價取券商賣出股數的最大值
top_sell = (
    merged.loc[merged["最高賣出"], ["股價", "賣出股數"]]
    .groupby("股價")["賣出股數"]
    .max()
)
price_summary["最高賣出股數"] = price_summary["股價"].map(top_sell).fillna(0)

# ====== 6) 輸出到 Excel ======
with ExcelWriter(output_path, engine="xlsxwriter") as writer:
    df.to_excel(writer, sheet_name="原始資料", index=False)
    price_summary.to_excel(writer, sheet_name="買賣價量與家數", index=False)
    merged.to_excel(writer, sheet_name="券商明細", index=False)

    workbook  = writer.book
    worksheet = writer.sheets["券商明細"]

    red_fmt    = workbook.add_format({"bg_color": "#FF9999"})
    green_fmt  = workbook.add_format({"bg_color": "#99FF99"})
    yellow_fmt = workbook.add_format({"bg_color": "#FFFF99"})

    broker_col_idx = merged.columns.get_loc("券商")
    buy_col_idx    = merged.columns.get_loc("買進股數")
    sell_col_idx   = merged.columns.get_loc("賣出股數")
    buy_ratio_idx  = merged.columns.get_loc("買進占比")
    sell_ratio_idx = merged.columns.get_loc("賣出占比")

    for row_num, row in enumerate(merged.itertuples(), start=1):
        if row.最高買進:
            worksheet.write(row_num, buy_col_idx, row.買進股數, red_fmt)
            worksheet.write(row_num, buy_ratio_idx, row.買進占比, red_fmt)
            worksheet.write(row_num, broker_col_idx, row.券商, yellow_fmt)
            highest_buy_col_idx = merged.columns.get_loc("最高買進")
            worksheet.write(row_num, highest_buy_col_idx, str(row.最高買進), red_fmt)

        if row.最高賣出:
            worksheet.write(row_num, sell_col_idx, row.賣出股數, green_fmt)
            worksheet.write(row_num, sell_ratio_idx, row.賣出占比, green_fmt)
            worksheet.write(row_num, broker_col_idx, row.券商, yellow_fmt)
            highest_sell_col_idx = merged.columns.get_loc("最高賣出")
            worksheet.write(row_num, highest_sell_col_idx, str(row.最高賣出), green_fmt)

print(f"已輸出到: {output_path}")

# ====== 7) Matplotlib 繪圖並直接插入 Excel ======
wb = load_workbook(output_path)
ws = wb.create_sheet("圖表")

x = np.arange(len(price_summary["股價"]))
width = 0.35

# 買進/賣出股數圖表
fig, ax = plt.subplots(figsize=(16,6))
ax.bar(x - width/2, price_summary["買進股數"], width, label="買進股數", color="#FFCCCC")
ax.bar(x - width/2, price_summary["最高買進股數"], width, label="最高買進股數", color="#8B0000")
ax.bar(x + width/2, price_summary["賣出股數"], width, label="賣出股數", color="#CCFFCC")
ax.bar(x + width/2, price_summary["最高賣出股數"], width, label="最高賣出股數", color="#006400")
ax.set_xticks(x)
ax.set_xticklabels(price_summary["股價"], rotation=90)
ax.set_xlabel("股價")
ax.set_ylabel("股數")
ax.set_title("各價位買進/賣出股數（含最高券商疊加）")
ax.legend()
# Y 軸細分化 + 虛線
y_max = int(ax.get_ylim()[1])
step = max(1, y_max // 20)
yticks = range(0, y_max + step, step)
ax.set_yticks(yticks)

for y in np.linspace(0, y_max, 20):
    ax.axhline(y, color="gray", linestyle="--", linewidth=0.5, alpha=0.5)

plt.tight_layout()
img_data1 = io.BytesIO()
fig.savefig(img_data1, format="png")
plt.close(fig)
img_data1.seek(0)
ws.add_image(Image(img_data1), "B2")

# 買進/賣出家數圖表
fig, ax = plt.subplots(figsize=(16,6))
ax.bar(x - width/2, price_summary["買進家數"], width, label="買進家數", color="red")
ax.bar(x + width/2, price_summary["賣出家數"], width, label="賣出家數", color="green")
ax.set_xticks(x)
ax.set_xticklabels(price_summary["股價"], rotation=90)
ax.set_xlabel("股價")
ax.set_ylabel("家數")
ax.set_title("各價位買進/賣出家數")
ax.legend()
# Y 軸細分化 + 虛線
y_max = int(ax.get_ylim()[1])
step = max(1, y_max // 20)
yticks = range(0, y_max + step, step)
ax.set_yticks(yticks)

for y in np.linspace(0, y_max, 20):
    ax.axhline(y, color="gray", linestyle="--", linewidth=0.5, alpha=0.5)

plt.tight_layout()

img_data2 = io.BytesIO()
fig.savefig(img_data2, format="png")
plt.close(fig)
img_data2.seek(0)
ws.add_image(Image(img_data2), "B35")

wb.save(output_path)

